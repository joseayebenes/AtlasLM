"""The traceability matrix as files for a review: CSV and XLSX for a spreadsheet, Markdown
to read or paste, JSON for a script (rows, summary, gaps and meta).

    write_matrix(res, prefix, formats, delimiter)   res = KnowledgeGraph.matrix(...)

The CSV is written for Excel: UTF-8 with a BOM (without it Excel reads the accents as
another encoding) and ";" between columns by default (the list separator of Excel in
Spanish; with "," everything lands in one column). The lists of a cell (code, tests,
documents) go one per line. The columns are those of MATRIX_COLUMNS with the DOORS attributes
of [sources.doors] attributes after the text (those of the matrix `meta`), and the CSV ends
with `exported_at`, the date of the DOORS export of the module of each row. The Markdown
starts with where the matrix comes from (matrix_header). XLSX needs openpyxl (the [xlsx] extra): a sheet with the matrix
(filters, widths, first row fixed), one with the gaps and one with the summary and the date
of the DOORS part of the graph.
"""
from __future__ import annotations

import csv
import io
import json
from datetime import datetime
from pathlib import Path

from atlas.config import MATRIX_FORMATS
from atlas.pipeline import atomic_write

__all__ = ["MATRIX_COLUMNS", "columns", "matrix_header", "summary_line", "write_matrix"]

MATRIX_COLUMNS = ["id", "text", "status", "verification", "derived", "code", "tests", "covering_tests", "procedures",
                  "documents"]
WIDTHS = {"id": 14, "text": 60, "status": 12, "verification": 14, "derived": 16, "code": 32, "tests": 30,
          "covering_tests": 30, "procedures": 18, "documents": 40}
GAPS = (("without_code", "Without code"), ("without_tests", "Without tests"),
        ("without_test_procedure", "Without a test procedure"))


def columns(res: dict) -> list[str]:
    """The columns of a matrix: those of MATRIX_COLUMNS, with the DOORS attributes of its
    `meta` in the place of status and verification."""
    attributes = (res.get("meta") or {}).get("attributes", ["status", "verification"])
    return ["id", "text", *attributes, *[c for c in MATRIX_COLUMNS[4:]]]


def _exported(res: dict, modules: dict[str, str]) -> dict[str, str]:
    """id -> the date of the DOORS export of its module (`modules`: id -> module)."""
    dates = {m: x.get("exported_at", "") for m, x in ((res.get("meta") or {}).get("modules") or {}).items()}
    return {r: dates.get(m, "") for r, m in modules.items()}


def matrix_header(meta: dict) -> list[str]:
    """Where a matrix comes from, as Markdown lines: the DOORS export of each module, the
    version of atlas, the builds of the graph and the revision of the code."""
    from atlas.doors import days_ago, export_age
    lines = []
    for m, x in (meta.get("modules") or {}).items():
        ago = days_ago(export_age(x["exported_at"]))
        lines.append(f"- DOORS {m}: " + (f"exported {x['exported_at'][:10]} ({ago})" if x["exported_at"]
                                          else "export date unknown") + (f", from {x['source']}" if x["source"] else ""))
    built = ", ".join(f"{origin} {at[:19]}" for origin, at in meta.get("built_at", {}).items())
    if meta:
        lines.append(f"- atlas {meta.get('atlas_version', '?')}" + (f"; graph built: {built}" if built else ""))
    if meta.get("code_revision"):
        lines.append(f"- code revision: {meta['code_revision']}")
    return lines


def _cell(value, separator: str) -> str:
    return separator.join(value) if isinstance(value, list) else str(value)


def _csv(res: dict, delimiter: str, exported: dict[str, str]) -> str:
    out = io.StringIO()
    w = csv.writer(out, delimiter=delimiter)
    cols = columns(res)
    w.writerow(cols + ["exported_at"])          # the DOORS export of the row's module
    w.writerows([_cell(row.get(c, ""), "\n") for c in cols] + [exported.get(row["id"], "")] for row in res["rows"])
    return out.getvalue()


def _markdown(res: dict) -> str:
    gaps = res["gaps"]
    cell = lambda v: _cell(v, "; ").replace("|", "\\|").replace("\n", " ")
    cols, header = columns(res), matrix_header(res.get("meta") or {})
    md = [f"# Traceability matrix: {res['module']}", "", *header, *([""] if header else []), summary_line(res), "",
          "| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
    md += ["| " + " | ".join(cell(row.get(c, "")) for c in cols) + " |" for row in res["rows"]]
    for key, title in GAPS:
        if key not in gaps:
            continue
        md += ["", f"## {title} ({gaps[key]['count']})", "", ", ".join(gaps[key]["items"]) or "-"]
    return "\n".join(md) + "\n"


def summary_line(res: dict) -> str:
    """'40 requirements: 4 with code, 1 with tests, 0 with test procedures, 40 in documents.'"""
    sm = res["summary"]
    procedures = f"{sm['with_procedures']} with test procedures, " if "with_procedures" in sm else ""
    return (f"{sm['requirements']} requirements: {sm['with_code']} with code, {sm['with_tests']} with tests, "
            f"{procedures}{sm['with_documents']} in documents.")


def _text(ws, row: int, column: int, value):
    """A cell whose text is never taken for a formula (a requirement that starts with '=')."""
    c = ws.cell(row=row, column=column, value=value)
    if isinstance(value, str) and value.startswith("="):
        c.data_type = "s"
    return c


def _xlsx(res: dict, path: Path, builds: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from atlas import __version__
    bold, wrap = Font(bold=True), Alignment(wrap_text=True, vertical="top")
    wb = Workbook()
    ws = wb.active
    ws.title = "Matrix"
    cols = columns(res)
    for j, name in enumerate(cols, 1):
        _text(ws, 1, j, name).font = bold
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS.get(name, 14)
    for i, row in enumerate(res["rows"], 2):
        for j, name in enumerate(cols, 1):
            _text(ws, i, j, _cell(row.get(name, ""), "\n")).alignment = wrap
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(res['rows']) + 1}"

    gaps = wb.create_sheet("Gaps")
    for j, name in enumerate(("gap", "requirement"), 1):
        _text(gaps, 1, j, name).font = bold
    gaps.column_dimensions["A"].width, gaps.column_dimensions["B"].width = 18, 18
    i = 2
    for key, title in GAPS:
        part = res["gaps"].get(key)
        if part is None:
            continue
        for item in part["items"]:
            _text(gaps, i, 1, title)
            _text(gaps, i, 2, item)
            i += 1
        if part.get("truncated"):
            _text(gaps, i, 1, title)
            _text(gaps, i, 2, f"… and {part['count'] - len(part['items'])} more")
            i += 1
    gaps.freeze_panes = "A2"
    gaps.auto_filter.ref = f"A1:B{max(i - 1, 1)}"

    summary = wb.create_sheet("Summary")
    sm, doors = res["summary"], builds.get("doors") or {}
    rows = [("Module", res["module"]), ("Requirements", sm["requirements"]), ("With code", sm["with_code"]),
            ("With tests", sm["with_tests"]), ("With test procedures", sm.get("with_procedures", 0)),
            ("In documents", sm["with_documents"]),
            ("Without code", res["gaps"]["without_code"]["count"]),
            ("Without tests", res["gaps"]["without_tests"]["count"]),
            ("Without a test procedure", (res["gaps"].get("without_test_procedure") or {}).get("count", 0)),
            ("DOORS read into the graph", (doors.get("built_at") or "-")[:19].replace("T", " ")),
            ("Written", datetime.now().isoformat(sep=" ", timespec="seconds")), ("atlas", __version__)]
    for i, (name, value) in enumerate(rows, 1):
        _text(summary, i, 1, name).font = bold
        _text(summary, i, 2, value)
    summary.column_dimensions["A"].width, summary.column_dimensions["B"].width = 28, 24
    tmp = path.with_name(path.name + ".tmp")
    wb.save(tmp)
    tmp.replace(path)


def write_matrix(res: dict, prefix: str | Path, formats=("csv", "md"), delimiter: str = ";",
                 builds: dict | None = None, modules: dict[str, str] | None = None) -> list[Path]:
    """Writes the matrix `res` (KnowledgeGraph.matrix) as <prefix>.csv, .md and .xlsx, the
    ones of `formats`, and returns their paths. `delimiter` separates the CSV columns;
    `builds` (GraphStore.builds()) gives the summary sheet of the XLSX the date of DOORS;
    `modules` (requirement -> its DOORS module) gives each CSV row the date of its export.
    "xlsx" without openpyxl raises ImportError before writing anything; a file open in Excel
    on Windows raises PermissionError."""
    unknown = sorted(set(formats) - set(MATRIX_FORMATS))
    if unknown:
        raise ValueError(f"unknown matrix formats {unknown} (valid: {', '.join(MATRIX_FORMATS)})")
    if len(delimiter) != 1:
        raise ValueError("the CSV delimiter must be one character")
    if "xlsx" in formats:
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            raise ImportError('"xlsx" in [graph] matrix_formats needs openpyxl: pip install "atlaslm[xlsx]"') from None
    prefix = Path(prefix)
    prefix.parent.mkdir(parents=True, exist_ok=True)
    written = []
    for kind in MATRIX_FORMATS:
        if kind not in formats:
            continue
        path = prefix.with_name(f"{prefix.name}.{kind}")
        if kind == "csv":
            atomic_write(path, _csv(res, delimiter, _exported(res, modules or {})).encode("utf-8-sig"))
        elif kind == "json":
            atomic_write(path, json.dumps({k: v for k, v in res.items() if k not in ("offset", "truncated")},
                                          ensure_ascii=False, indent=1, default=str))
        elif kind == "md":
            atomic_write(path, _markdown(res))
        else:
            _xlsx(res, path, builds or {})
        written.append(path)
    return written
